"""Generate Word doc for scorecard audit: questions, scoring, logic, recommendations."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(6)

# ── Title ──
title = doc.add_heading('Broker Tax Scorecard — Audit Document', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph(
    'This document contains all questions, scoring values, grading logic, '
    'snapshot rows, and recommendation text from the scorecard. '
    'Edit as needed and return — changes will be applied to the code.'
)
doc.add_paragraph('')

# ═══════════════════════════════════════════════════════════════
# SECTION 1: QUESTIONS
# ═══════════════════════════════════════════════════════════════
doc.add_heading('1. Questions', level=1)
doc.add_paragraph(
    'Question order: Income → Income Type → Entity (conditional) → Retirement '
    '→ Estimated Tax → Expenses → Planning → S-Corp Salary (conditional)'
)

questions = [
    {
        'id': 'income',
        'num': 1,
        'text': "What's your annual net income?",
        'condition': 'Always shown',
        'options': [
            ('Under $150K', 'under150'),
            ('$150K – $300K', '150-300'),
            ('$300K – $500K', '300-500'),
            ('$500K – $1M', '500-1m'),
            ('Over $1M', 'over1m'),
        ],
    },
    {
        'id': 'income_type',
        'num': 2,
        'text': 'How do you receive your commission income?',
        'condition': 'Always shown',
        'options': [
            ('1099 independent contractor', '1099'),
            ('Mix of W-2 and 1099', 'mix'),
            ("I'm a W-2 employee", 'w2'),
        ],
    },
    {
        'id': 'entity',
        'num': 3,
        'text': 'How is your business currently structured?',
        'condition': 'Only shown if income_type ≠ "w2"',
        'options': [
            ('Sole Proprietor / Schedule C', 'sole_prop'),
            ('Single-Member LLC (taxed as sole prop)', 'smllc'),
            ('S-Corporation', 's_corp'),
            ("I'm not sure", 'not_sure'),
        ],
    },
    {
        'id': 'retirement',
        'num': 4,
        'text': 'Do you have a retirement plan set up through your business?',
        'condition': 'Always shown',
        'options': [
            ('Yes, and I max it out', 'max'),
            ("Yes, but I'm not contributing much", 'low'),
            ('No retirement plan through my business', 'none'),
            ('I only use a personal IRA', 'ira'),
        ],
    },
    {
        'id': 'estimated_tax',
        'num': 5,
        'text': 'How are you handling your tax payments throughout the year?',
        'condition': 'Always shown',
        'options': [
            ("I project my taxes annually and pay them in December through my S Corp's paycheck", 'w2_withholding'),
            ('I consistently make quarterly payments', 'quarterly'),
            ('I make payments sometimes, but not regularly', 'sometimes'),
            ('I just pay what my CPA says upon filing my return', 'no'),
            ("I'm a W-2 employee so taxes are withheld automatically", 'w2_auto'),
        ],
    },
    {
        'id': 'expenses',
        'num': 6,
        'text': 'How well are you tracking business deductions?',
        'condition': 'Always shown',
        'options': [
            ('Detailed system — I track everything', 'detailed'),
            ('Loosely — I keep some receipts', 'loosely'),
            ("I'm not really tracking deductions", 'not_tracking'),
        ],
    },
    {
        'id': 'planning',
        'num': 7,
        'text': 'When was your last proactive tax planning session (not just filing your return)?',
        'condition': 'Always shown',
        'options': [
            ('Within the last 6 months', '6months'),
            ('Within the last year', 'year'),
            ('Over a year ago', 'over_year'),
            ('Never had one', 'never'),
        ],
    },
    {
        'id': 'scorp_salary',
        'num': 8,
        'text': 'Roughly what percentage of your S-Corp net income do you pay yourself as W-2 salary?',
        'condition': 'Only shown if entity = "s_corp"',
        'options': [
            ('Less than 30%', 'under30'),
            ('30% – 50%', '30-50'),
            ('Over 50%', 'over50'),
            ("I'm not sure", 'not_sure'),
        ],
    },
]

for q in questions:
    doc.add_heading(f"Q{q['num']}: {q['id']}", level=2)
    doc.add_paragraph(f"Question text:  {q['text']}")
    doc.add_paragraph(f"Condition:  {q['condition']}")
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    hdr[0].text = 'Option Label'
    hdr[1].text = 'Value (code key)'
    for label, val in q['options']:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = val
    doc.add_paragraph('')

# ═══════════════════════════════════════════════════════════════
# SECTION 2: SCORING
# ═══════════════════════════════════════════════════════════════
doc.add_heading('2. Scoring', level=1)

doc.add_heading('Derived flags (not scored directly)', level=2)
flags = [
    ('isW2Employee', 'income_type === "w2"'),
    ('isSCorp', 'entity === "s_corp"'),
    ('showedEntity', '!isW2Employee  (entity question was shown)'),
    ('incomeHigh', 'income is one of: 150-300, 300-500, 500-1m, over1m'),
    ('income300Plus', 'income is one of: 300-500, 500-1m, over1m'),
]
table = doc.add_table(rows=1, cols=2)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
hdr[0].text = 'Flag'
hdr[1].text = 'Condition'
for flag, cond in flags:
    row = table.add_row().cells
    row[0].text = flag
    row[1].text = cond
doc.add_paragraph('')

doc.add_heading('Income Type scoring (always scored)', level=2)
table = doc.add_table(rows=1, cols=2)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
hdr[0].text = 'Answer'
hdr[1].text = 'Points'
for ans, pts in [('1099', '10'), ('mix', '6'), ('w2', '4')]:
    row = table.add_row().cells
    row[0].text = ans
    row[1].text = pts
doc.add_paragraph('')

doc.add_heading('Entity scoring (only if showedEntity = true)', level=2)
table = doc.add_table(rows=1, cols=2)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
hdr[0].text = 'Answer'
hdr[1].text = 'Points'
for ans, pts in [('s_corp', '10'), ('smllc', '3'), ('sole_prop', '2'), ('not_sure', '1')]:
    row = table.add_row().cells
    row[0].text = ans
    row[1].text = pts
doc.add_paragraph('')

doc.add_heading('Penalty', level=2)
doc.add_paragraph('−5 points if incomeHigh AND not S-Corp AND not W2 employee')
doc.add_paragraph('')

doc.add_heading('Retirement scoring', level=2)
table = doc.add_table(rows=1, cols=2)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
hdr[0].text = 'Answer'
hdr[1].text = 'Points'
for ans, pts in [('max', '10'), ('low', '5'), ('ira', '3'), ('none', '0')]:
    row = table.add_row().cells
    row[0].text = ans
    row[1].text = pts
doc.add_paragraph('')

doc.add_heading('Estimated Tax scoring', level=2)
table = doc.add_table(rows=1, cols=2)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
hdr[0].text = 'Answer'
hdr[1].text = 'Points'
for ans, pts in [('w2_withholding', '10'), ('quarterly', '8'), ('w2_auto', '7'), ('sometimes', '4'), ('no', '0')]:
    row = table.add_row().cells
    row[0].text = ans
    row[1].text = pts
doc.add_paragraph('')

doc.add_heading('Expenses scoring', level=2)
table = doc.add_table(rows=1, cols=2)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
hdr[0].text = 'Answer'
hdr[1].text = 'Points'
for ans, pts in [('detailed', '10'), ('loosely', '4'), ('not_tracking', '0')]:
    row = table.add_row().cells
    row[0].text = ans
    row[1].text = pts
doc.add_paragraph('')

doc.add_heading('Planning scoring', level=2)
table = doc.add_table(rows=1, cols=2)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
hdr[0].text = 'Answer'
hdr[1].text = 'Points'
for ans, pts in [('6months', '10'), ('year', '6'), ('over_year', '2'), ('never', '0')]:
    row = table.add_row().cells
    row[0].text = ans
    row[1].text = pts
doc.add_paragraph('')

doc.add_heading('S-Corp Salary scoring (only if isSCorp)', level=2)
table = doc.add_table(rows=1, cols=2)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
hdr[0].text = 'Answer'
hdr[1].text = 'Points'
for ans, pts in [('30-50', '10'), ('over50', '5'), ('under30', '2'), ('not_sure', '1')]:
    row = table.add_row().cells
    row[0].text = ans
    row[1].text = pts
doc.add_paragraph('')

# ═══════════════════════════════════════════════════════════════
# SECTION 3: MAX SCORE & GRADING
# ═══════════════════════════════════════════════════════════════
doc.add_heading('3. Max Score & Grading', level=1)

doc.add_heading('Max score calculation', level=2)
doc.add_paragraph('Base max = 50  (income_type + retirement + estimated_tax + expenses + planning)')
doc.add_paragraph('+10 if showedEntity (entity question was shown)')
doc.add_paragraph('+10 if isSCorp (scorp_salary question was shown)')
doc.add_paragraph('Percentage = total / maxScore × 100')
doc.add_paragraph('')

doc.add_heading('Grade thresholds', level=2)
table = doc.add_table(rows=1, cols=4)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
hdr[0].text = 'Grade'
hdr[1].text = 'Threshold'
hdr[2].text = 'Label'
hdr[3].text = 'Summary Text'
grades = [
    ('A', '≥ 85%', 'Well Optimized', "Your tax structure is well optimized. You're ahead of most brokers."),
    ('B', '≥ 70%', 'Good Foundation', "You've got a solid foundation — a few tweaks could save you more."),
    ('C', '≥ 55%', 'Room to Improve', "You're likely leaving real money on the table. Let's fix that."),
    ('D', '≥ 40%', 'Leaving Money on the Table', "There are significant gaps in your tax strategy."),
    ('F', '< 40%', 'Needs Attention', "Your tax structure needs attention — but the good news is the fixes are straightforward."),
]
for g, thresh, label, summary in grades:
    row = table.add_row().cells
    row[0].text = g
    row[1].text = thresh
    row[2].text = label
    row[3].text = summary
doc.add_paragraph('')

doc.add_heading('Critical flag', level=2)
doc.add_paragraph('criticalRetirement = income is $300K+ AND retirement answer is "none"')
doc.add_paragraph('')

# ═══════════════════════════════════════════════════════════════
# SECTION 4: SNAPSHOT ROWS  (Current → Optimal)
# ═══════════════════════════════════════════════════════════════
doc.add_heading('4. Snapshot Rows (Current → Optimal)', level=1)
doc.add_paragraph('These appear in the results as a "where you are vs. where you could be" comparison.')
doc.add_paragraph('')

snapshots = [
    ('Entity gap', 'entity ≠ s_corp AND not W2 AND incomeHigh',
     [('Schedule C / Sole Prop  OR  Single-Member LLC  OR  Unsure of entity structure',
       'S-Corporation saving $15K–$40K+/yr in SE tax')]),
    ('S-Corp salary too low', 'isSCorp AND scorp_salary = "under30"',
     [('S-Corp salary too low (<30%)',
       'Right-sized reasonable salary — max savings, IRS-compliant')]),
    ('S-Corp salary too high', 'isSCorp AND scorp_salary = "over50"',
     [('S-Corp salary too high (>50%)',
       'Optimized salary-to-distribution ratio saving thousands in FICA')]),
    ('S-Corp salary not sure', 'isSCorp AND scorp_salary = "not_sure"',
     [('Unsure of salary-to-distribution ratio',
       'Dialed-in ratio maximizing savings and compliance')]),
    ('No retirement plan', 'retirement = "none" OR "ira"',
     [('No business retirement plan',
       'Solo 401(k) sheltering up to $69K/yr tax-deferred')]),
    ('Low retirement', 'retirement = "low"',
     [('Low retirement contributions',
       'Maxed-out retirement plan sheltering up to $69K/yr')]),
    ('Quarterly payments (S-Corp)', 'estimated_tax = "quarterly" AND isSCorp',
     [('Quarterly estimated payments',
       'Year-end W-2 withholding strategy — no underpayment penalties')]),
    ('Sporadic payments', 'estimated_tax = "sometimes"',
     [('Sporadic estimated payments',
       'Structured payments avoiding underpayment penalties')]),
    ('Paying at filing only', 'estimated_tax = "no"',
     [('Paying only at filing time',
       'Structured payments avoiding underpayment penalties')]),
    ('Loose expense tracking', 'expenses = "loosely"',
     [('Sporadic expense tracking',
       'Maximized deductions with a clean system')]),
    ('No expense tracking', 'expenses = "not_tracking"',
     [('No expense tracking',
       'Maximized deductions with a clean system')]),
    ('No recent planning', 'planning = "over_year" OR "never"',
     [('No recent tax planning',
       'Year-round strategy capturing every opportunity')]),
]

for title_text, condition, rows_data in snapshots:
    doc.add_heading(title_text, level=2)
    doc.add_paragraph(f"Condition:  {condition}")
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Light Grid Accent 1'
    hdr = table.rows[0].cells
    hdr[0].text = 'Current'
    hdr[1].text = 'Optimal'
    for current, optimal in rows_data:
        row = table.add_row().cells
        row[0].text = current
        row[1].text = optimal
    doc.add_paragraph('')

# ═══════════════════════════════════════════════════════════════
# SECTION 5: RECOMMENDATIONS
# ═══════════════════════════════════════════════════════════════
doc.add_heading('5. Recommendations', level=1)
doc.add_paragraph('Up to 3 recommendations are shown, selected in priority order from top to bottom.')
doc.add_paragraph('')

recs = [
    (
        '1. S-Corp gap',
        'entity ≠ s_corp AND not W2 employee AND incomeHigh',
        "You're earning [income bracket] on a [entity label] — an S-Corp election could reduce your "
        "self-employment taxes by $15,000 to $40,000+ per year. This is likely your single biggest savings opportunity."
    ),
    (
        '2. S-Corp salary too low',
        'isSCorp AND scorp_salary = "under30"',
        "Your S-Corp salary may be set too low relative to your income. The IRS requires a 'reasonable salary' "
        "— setting it below ~30% of net income is a common audit trigger. Getting this number right protects you "
        "while still maximizing your tax savings."
    ),
    (
        '3. S-Corp salary too high',
        'isSCorp AND scorp_salary = "over50"',
        "You may be paying yourself more salary than necessary through your S-Corp, which means you're overpaying "
        "FICA taxes. Dialing this back to the right level could put thousands back in your pocket."
    ),
    (
        '4. S-Corp salary not sure',
        'isSCorp AND scorp_salary = "not_sure"',
        "Not knowing your salary-to-distribution ratio is a red flag — this is one of the most important numbers "
        "in your S-Corp. Getting this dialed in is critical for both tax savings and IRS compliance."
    ),
    (
        '5. Retirement gap',
        'retirement = "none" OR "ira"',
        "Without a business retirement plan, you're missing out on one of the best tax shelters available. "
        "A Solo 401(k) could let you defer up to $69,000/year in pre-tax income."
    ),
    (
        '6. No tax planning',
        'planning = "never" OR "over_year"',
        "Most of the biggest tax savings require planning before year-end — not at filing time. A proactive "
        "tax planning session could surface opportunities you're currently missing."
    ),
    (
        '7. Expense tracking gap',
        'expenses = "loosely" OR "not_tracking"',
        "Tightening up your deduction tracking could put thousands back in your pocket. Many brokers under-deduct "
        "because they don't have a system in place."
    ),
    (
        '8. Paying at filing only',
        'estimated_tax = "no" AND not W2 employee',
        "Paying only at filing time can lead to IRS underpayment penalties. Getting on a structured payment "
        "schedule protects you and improves cash flow predictability."
    ),
    (
        '9. Not using W-2 withholding strategy',
        'isSCorp AND estimated_tax = "quarterly"',
        "If you're an S-Corp, there's a smarter way to handle tax payments: withhold through your W-2 salary "
        "with a year-end true-up. Unlike quarterly estimates, W-2 withholding is treated as paid evenly "
        "throughout the year — which means no underpayment penalties, even if you do it all in December."
    ),
    (
        '10. W-2 employee general',
        'isW2Employee',
        "Being W-2 limits some optimization strategies, but there may still be opportunities worth exploring "
        "— especially around retirement plans and deduction strategies."
    ),
]

for title_text, condition, text in recs:
    doc.add_heading(title_text, level=2)
    p = doc.add_paragraph()
    p.add_run('Condition:  ').bold = True
    p.add_run(condition)
    doc.add_paragraph('')
    p2 = doc.add_paragraph()
    p2.add_run('Text shown to user:').bold = True
    doc.add_paragraph(text)
    doc.add_paragraph('')

# ── Save ──
output_path = '/home/user/claude-code/Scorecard_Audit.docx'
doc.save(output_path)
print(f'Saved to {output_path}')
